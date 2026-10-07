"""Reporting engine.

Every report is produced as a table (``columns`` + ``rows`` + ``totals``) so a
single export layer can render CSV, Excel or PDF and the Flutter client can show
the same data in a data table with filters, grouping and sorting.
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any, Callable, Iterable, Sequence

from sqlalchemy import Select, and_, func, or_, select, text
from sqlalchemy.orm import Session

from app.core.enums import AccountType, DocumentStatus, MovementType, PartyType
from app.core.coercion import as_uuid
from app.core.errors import BusinessRuleError, ValidationFailure
from app.models.accounting import (
    Account,
    CustomerLedgerEntry,
    JournalEntry,
    JournalEntryLine,
    SupplierLedgerEntry,
)
from app.models.hr import AttendanceRecord, Employee, LeaveRequest, Payslip, PayrollRun
from app.models.inventory import StockBalance, StockLedgerEntry
from app.models.masterdata import Customer, Product, Supplier, Warehouse
from app.models.platform import Company
from app.models.projects import Project, TimesheetLine
from app.models.purchasing import PurchaseInvoice, PurchaseOrder, SupplierQuotation
from app.models.sales import CreditNote, SalesInvoice, SalesInvoiceLine, SalesOrder
from app.models.treasury import BankAccount, CashAccount, Payment
from app.services.posting_service import NORMAL_BALANCE_FOR_TYPE, PostingService, money, quantity

Column = dict[str, Any]


class ReportResult:
    """Uniform report payload."""

    def __init__(
        self,
        *,
        code: str,
        title: str,
        columns: Sequence[Column],
        rows: Sequence[dict[str, Any]],
        totals: dict[str, Any] | None = None,
        meta: dict[str, Any] | None = None,
    ) -> None:
        self.code = code
        self.title = title
        self.columns = list(columns)
        self.rows = [dict(row) for row in rows]
        self.totals = totals or {}
        self.meta = meta or {}

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "title": self.title,
            "columns": self.columns,
            "rows": [
                {key: (str(value) if isinstance(value, (Decimal, date, datetime, uuid.UUID)) else value) for key, value in row.items()}
                for row in self.rows
            ],
            "totals": {key: str(value) if isinstance(value, Decimal) else value for key, value in self.totals.items()},
            "meta": self.meta,
            "generated_at": datetime.now(UTC).isoformat(),
        }

    # ------------------------------------------------------------- exports
    def to_csv(self) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([column.get("title", column["key"]) for column in self.columns])
        for row in self.rows:
            writer.writerow([_cell(row.get(column["key"])) for column in self.columns])
        if self.totals:
            writer.writerow([])
            for key, value in self.totals.items():
                writer.writerow([key, _cell(value)])
        return buffer.getvalue().encode("utf-8-sig")

    def to_excel(self) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:  # pragma: no cover
            raise BusinessRuleError("Excel export requires the openpyxl package") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Report"
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        sheet.append([self.title])
        sheet.append([f"Generated at {datetime.now(UTC).isoformat(timespec='seconds')}"])
        sheet.append([])
        sheet.append([column.get("title", column["key"]) for column in self.columns])
        for cell in sheet[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in self.rows:
            sheet.append([_cell(row.get(column["key"])) for column in self.columns])
        if self.totals:
            sheet.append([])
            for key, value in self.totals.items():
                sheet.append([key, _cell(value)])
        for index, column in enumerate(self.columns, start=1):
            width = max(len(str(column.get("title", column["key"]))) + 4, 14)
            sheet.column_dimensions[sheet.cell(row=4, column=index).column_letter].width = width
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def to_pdf(self) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:  # pragma: no cover
            raise BusinessRuleError("PDF export requires the reportlab package") from exc

        styles = getSampleStyleSheet()
        stream = io.BytesIO()
        document = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm)
        elements: list[Any] = [
            Paragraph(self.title, styles["Title"]),
            Paragraph(f"Generated {datetime.now(UTC).strftime('%Y-%m-%d %H:%M')} UTC", styles["Normal"]),
            Spacer(1, 6),
        ]
        header = [column.get("title", column["key"]) for column in self.columns]
        data = [header] + [[_cell(row.get(column["key"])) for column in self.columns] for row in self.rows[:2000]]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
                ]
            )
        )
        elements.append(table)
        if self.totals:
            elements.append(Spacer(1, 8))
            totals_table = Table([[key, _cell(value)] for key, value in self.totals.items()])
            totals_table.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (0, 0), (-1, -1), "RIGHT")]))
            elements.append(totals_table)
        document.build(elements)
        return stream.getvalue()


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10] if isinstance(value, date) and not isinstance(value, datetime) else value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


class ReportService:
    def __init__(self, db: Session, company_id: uuid.UUID) -> None:
        self.db = db
        self.company_id = company_id
        self.posting = PostingService(db, company_id)

    # ------------------------------------------------------------- financial
    def trial_balance(
        self,
        *,
        date_from: date | None = None,
        date_to: date | None = None,
        include_zero: bool = False,
        branch_id: uuid.UUID | None = None,
    ) -> ReportResult:
        accounts = list(
            self.db.execute(
                select(Account)
                .where(Account.company_id == self.company_id, Account.is_active.is_(True))
                .order_by(Account.code)
            ).scalars().all()
        )
        rows: list[dict[str, Any]] = []
        total_debit = Decimal("0")
        total_credit = Decimal("0")
        for account in accounts:
            debit, credit = self.posting.account_balance(
                account.id, date_from=date_from, date_to=date_to, branch_id=branch_id
            )
            balance = debit - credit
            if not include_zero and debit == 0 and credit == 0:
                continue
            total_debit += debit
            total_credit += credit
            rows.append(
                {
                    "account_code": account.code,
                    "account_name": account.name,
                    "account_name_ar": account.name_ar,
                    "account_type": account.account_type,
                    "opening_balance": str(account.opening_balance or 0),
                    "debit": str(debit),
                    "credit": str(credit),
                    "balance": str(balance),
                    "balance_type": "debit" if balance >= 0 else "credit",
                }
            )
        return ReportResult(
            code="trial_balance",
            title="Trial Balance",
            columns=[
                {"key": "account_code", "title": "Code"},
                {"key": "account_name", "title": "Account"},
                {"key": "account_name_ar", "title": "الاسم"},
                {"key": "account_type", "title": "Type"},
                {"key": "debit", "title": "Debit", "type": "money"},
                {"key": "credit", "title": "Credit", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={
                "total_debit": str(money(total_debit)),
                "total_credit": str(money(total_credit)),
                "difference": str(money(total_debit - total_credit)),
            },
            meta={"date_from": date_from, "date_to": date_to},
        )

    def general_ledger(
        self,
        *,
        account_id: uuid.UUID | None = None,
        date_from: date | None = None,
        date_to: date | None = None,
        party_id: uuid.UUID | None = None,
        limit: int = 5000,
    ) -> ReportResult:
        stmt = (
            select(JournalEntryLine, JournalEntry, Account)
            .join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .join(Account, Account.id == JournalEntryLine.account_id)
            .where(
                JournalEntryLine.company_id == self.company_id,
                JournalEntry.status == DocumentStatus.POSTED.value,
                JournalEntry.deleted_at.is_(None),
            )
            .order_by(JournalEntry.entry_date, JournalEntry.entry_no, JournalEntryLine.sequence_no)
            .limit(limit)
        )
        if account_id:
            stmt = stmt.where(JournalEntryLine.account_id == account_id)
        if date_from:
            stmt = stmt.where(JournalEntry.entry_date >= date_from)
        if date_to:
            stmt = stmt.where(JournalEntry.entry_date <= date_to)
        if party_id:
            stmt = stmt.where(JournalEntryLine.party_id == party_id)

        rows: list[dict[str, Any]] = []
        running = Decimal("0")
        for line, entry, account in self.db.execute(stmt).all():
            amount = Decimal(line.debit or 0) - Decimal(line.credit or 0)
            if account.normal_balance != NORMAL_BALANCE_FOR_TYPE.get(account.account_type, "debit"):
                running -= amount
            else:
                running += amount
            rows.append(
                {
                    "date": entry.entry_date,
                    "entry_no": entry.entry_no,
                    "account_code": account.code,
                    "account_name": account.name,
                    "description": line.description or entry.description,
                    "reference": entry.source_document_no or entry.reference,
                    "debit": str(money(line.debit)),
                    "credit": str(money(line.credit)),
                    "balance": str(money(running)),
                    "party_id": line.party_id,
                }
            )
        return ReportResult(
            code="general_ledger",
            title="General Ledger",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "entry_no", "title": "Entry"},
                {"key": "account_code", "title": "Code"},
                {"key": "account_name", "title": "Account"},
                {"key": "description", "title": "Description"},
                {"key": "debit", "title": "Debit", "type": "money"},
                {"key": "credit", "title": "Credit", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={"closing_balance": str(money(running))},
            meta={"account_id": account_id, "date_from": date_from, "date_to": date_to},
        )

    def income_statement(
        self, *, date_from: date, date_to: date, branch_id: uuid.UUID | None = None, compare: bool = False
    ) -> ReportResult:
        rows, totals = self._account_type_rows(
            [AccountType.REVENUE, AccountType.EXPENSE], date_from=date_from, date_to=date_to, branch_id=branch_id
        )
        revenue = sum((Decimal(row["amount"]) for row in rows if row["account_type"] == AccountType.REVENUE.value), Decimal("0"))
        expense = sum((Decimal(row["amount"]) for row in rows if row["account_type"] == AccountType.EXPENSE.value), Decimal("0"))
        net = money(revenue - expense)
        result = ReportResult(
            code="income_statement",
            title="Profit & Loss",
            columns=[
                {"key": "account_code", "title": "Code"},
                {"key": "account_name", "title": "Account"},
                {"key": "account_type", "title": "Type"},
                {"key": "amount", "title": "Amount", "type": "money"},
            ],
            rows=rows,
            totals={
                "total_revenue": str(money(revenue)),
                "total_expense": str(money(expense)),
                "net_profit": str(net),
                "net_margin_percent": str(
                    (net / revenue * Decimal("100")).quantize(Decimal("0.01")) if revenue else Decimal("0")
                ),
            },
            meta={"date_from": date_from, "date_to": date_to},
        )
        if compare:
            period = (date_to - date_from) or timedelta(days=1)
            previous_from = date_from - period - timedelta(days=1)
            previous_to = date_from - timedelta(days=1)
            previous_rows, _ = self._account_type_rows(
                [AccountType.REVENUE, AccountType.EXPENSE],
                date_from=previous_from,
                date_to=previous_to,
                branch_id=branch_id,
            )
            previous_revenue = sum(
                (Decimal(row["amount"]) for row in previous_rows if row["account_type"] == AccountType.REVENUE.value),
                Decimal("0"),
            )
            previous_expense = sum(
                (Decimal(row["amount"]) for row in previous_rows if row["account_type"] == AccountType.EXPENSE.value),
                Decimal("0"),
            )
            result.meta["comparison"] = {
                "period": {"from": previous_from.isoformat(), "to": previous_to.isoformat()},
                "revenue": str(money(previous_revenue)),
                "expense": str(money(previous_expense)),
                "net_profit": str(money(previous_revenue - previous_expense)),
            }
        return result

    def balance_sheet(self, *, as_of: date, branch_id: uuid.UUID | None = None) -> ReportResult:
        rows, totals = self._account_type_rows(
            [AccountType.ASSET, AccountType.LIABILITY, AccountType.EQUITY],
            date_from=None,
            date_to=as_of,
            branch_id=branch_id,
        )
        revenue_rows, _ = self._account_type_rows(
            [AccountType.REVENUE, AccountType.EXPENSE], date_from=None, date_to=as_of, branch_id=branch_id
        )
        revenue = sum((Decimal(row["amount"]) for row in revenue_rows if row["account_type"] == AccountType.REVENUE.value), Decimal("0"))
        expense = sum((Decimal(row["amount"]) for row in revenue_rows if row["account_type"] == AccountType.EXPENSE.value), Decimal("0"))
        net_profit = money(revenue - expense)
        assets = sum((Decimal(row["amount"]) for row in rows if row["account_type"] == AccountType.ASSET.value), Decimal("0"))
        liabilities = sum((Decimal(row["amount"]) for row in rows if row["account_type"] == AccountType.LIABILITY.value), Decimal("0"))
        equity = sum((Decimal(row["amount"]) for row in rows if row["account_type"] == AccountType.EQUITY.value), Decimal("0"))
        return ReportResult(
            code="balance_sheet",
            title="Balance Sheet",
            columns=[
                {"key": "account_code", "title": "Code"},
                {"key": "account_name", "title": "Account"},
                {"key": "account_type", "title": "Type"},
                {"key": "amount", "title": "Amount", "type": "money"},
            ],
            rows=rows,
            totals={
                "total_assets": str(money(assets)),
                "total_liabilities": str(money(liabilities)),
                "total_equity": str(money(equity)),
                "current_result": str(net_profit),
                "total_liabilities_and_equity": str(money(liabilities + equity + net_profit)),
                "balanced": bool(abs(money(assets - (liabilities + equity + net_profit))) <= Decimal("0.01")),
            },
            meta={"as_of": as_of},
        )

    def cash_flow(self, *, date_from: date, date_to: date) -> ReportResult:
        stmt = (
            select(
                Account.cash_flow_category,
                func.coalesce(func.sum(JournalEntryLine.debit - JournalEntryLine.credit), 0),
            )
            .join(JournalEntryLine, JournalEntryLine.account_id == Account.id)
            .join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .where(
                JournalEntryLine.company_id == self.company_id,
                JournalEntry.status == DocumentStatus.POSTED.value,
                JournalEntry.entry_date >= date_from,
                JournalEntry.entry_date <= date_to,
                Account.is_cash_account.is_(True) | Account.is_bank_account.is_(True),
            )
            .group_by(Account.cash_flow_category)
        )
        rows = []
        total = Decimal("0")
        for category, amount in self.db.execute(stmt).all():
            rows.append(
                {
                    "category": category or "operating",
                    "inflow": str(money(amount if amount > 0 else 0)),
                    "outflow": str(money(-amount if amount < 0 else 0)),
                    "net": str(money(amount)),
                }
            )
            total += Decimal(amount or 0)
        opening = self._cash_opening(date_from)
        return ReportResult(
            code="cash_flow",
            title="Cash Flow Statement",
            columns=[
                {"key": "category", "title": "Activity"},
                {"key": "inflow", "title": "Inflow", "type": "money"},
                {"key": "outflow", "title": "Outflow", "type": "money"},
                {"key": "net", "title": "Net", "type": "money"},
            ],
            rows=rows,
            totals={
                "opening_cash": str(opening),
                "net_change": str(money(total)),
                "closing_cash": str(money(opening + total)),
            },
            meta={"date_from": date_from, "date_to": date_to},
        )

    def _cash_opening(self, date_from: date) -> Decimal:
        stmt = (
            select(func.coalesce(func.sum(JournalEntryLine.debit - JournalEntryLine.credit), 0))
            .join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .join(Account, Account.id == JournalEntryLine.account_id)
            .where(
                JournalEntryLine.company_id == self.company_id,
                JournalEntry.status == DocumentStatus.POSTED.value,
                JournalEntry.entry_date < date_from,
                Account.is_cash_account.is_(True) | Account.is_bank_account.is_(True),
            )
        )
        return money(self.db.execute(stmt).scalar_one())

    def _account_type_rows(
        self,
        account_types: Sequence[AccountType],
        *,
        date_from: date | None,
        date_to: date | None,
        branch_id: uuid.UUID | None = None,
    ) -> tuple[list[dict[str, Any]], dict[str, Decimal]]:
        types = [item.value for item in account_types]
        stmt = (
            select(
                Account.id,
                Account.code,
                Account.name,
                Account.name_ar,
                Account.account_type,
                Account.normal_balance,
                func.coalesce(func.sum(JournalEntryLine.debit), 0),
                func.coalesce(func.sum(JournalEntryLine.credit), 0),
            )
            .join(JournalEntryLine, JournalEntryLine.account_id == Account.id)
            .join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .where(
                Account.company_id == self.company_id,
                Account.account_type.in_(types),
                JournalEntry.status == DocumentStatus.POSTED.value,
                JournalEntry.deleted_at.is_(None),
            )
            .group_by(Account.id, Account.code, Account.name, Account.name_ar, Account.account_type, Account.normal_balance)
            .order_by(Account.code)
        )
        if date_from:
            stmt = stmt.where(JournalEntry.entry_date >= date_from)
        if date_to:
            stmt = stmt.where(JournalEntry.entry_date <= date_to)
        if branch_id:
            stmt = stmt.where(JournalEntryLine.branch_id == branch_id)
        rows: list[dict[str, Any]] = []
        for account_id, code, name, name_ar, account_type, normal_balance, debit, credit in self.db.execute(stmt).all():
            debit = money(debit)
            credit = money(credit)
            amount = (debit - credit) if normal_balance == "debit" else (credit - debit)
            rows.append(
                {
                    "account_id": account_id,
                    "account_code": code,
                    "account_name": name,
                    "account_name_ar": name_ar,
                    "account_type": account_type,
                    "debit": str(debit),
                    "credit": str(credit),
                    "amount": str(money(amount)),
                }
            )
        return rows, {}

    # ---------------------------------------------------------------- parties
    def customer_statement(self, *, customer_id: uuid.UUID, date_from: date, date_to: date) -> ReportResult:
        customer = self.db.get(Customer, customer_id)
        opening = self.db.execute(
            select(
                func.coalesce(func.sum(CustomerLedgerEntry.debit - CustomerLedgerEntry.credit), 0)
            ).where(
                CustomerLedgerEntry.company_id == self.company_id,
                CustomerLedgerEntry.customer_id == customer_id,
                CustomerLedgerEntry.document_date < date_from,
            )
        ).scalar_one()
        entries = self.db.execute(
            select(CustomerLedgerEntry)
            .where(
                CustomerLedgerEntry.company_id == self.company_id,
                CustomerLedgerEntry.customer_id == customer_id,
                CustomerLedgerEntry.document_date >= date_from,
                CustomerLedgerEntry.document_date <= date_to,
            )
            .order_by(CustomerLedgerEntry.document_date, CustomerLedgerEntry.created_at)
        ).scalars().all()
        running = money(opening)
        rows = []
        for entry in entries:
            running = money(running + Decimal(entry.debit or 0) - Decimal(entry.credit or 0))
            rows.append(
                {
                    "date": entry.document_date,
                    "document_type": entry.document_type,
                    "document_no": entry.document_no,
                    "due_date": entry.due_date,
                    "debit": str(money(entry.debit)),
                    "credit": str(money(entry.credit)),
                    "balance": str(running),
                }
            )
        return ReportResult(
            code="customer_statement",
            title=f"Customer Statement — {customer.name if customer else ''}",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "document_type", "title": "Type"},
                {"key": "document_no", "title": "Document"},
                {"key": "due_date", "title": "Due", "type": "date"},
                {"key": "debit", "title": "Debit", "type": "money"},
                {"key": "credit", "title": "Credit", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={"opening_balance": str(money(opening)), "closing_balance": str(running)},
            meta={"customer_id": customer_id, "date_from": date_from, "date_to": date_to},
        )

    def supplier_statement(self, *, supplier_id: uuid.UUID, date_from: date, date_to: date) -> ReportResult:
        supplier = self.db.get(Supplier, supplier_id)
        opening = self.db.execute(
            select(func.coalesce(func.sum(SupplierLedgerEntry.credit - SupplierLedgerEntry.debit), 0)).where(
                SupplierLedgerEntry.company_id == self.company_id,
                SupplierLedgerEntry.supplier_id == supplier_id,
                SupplierLedgerEntry.document_date < date_from,
            )
        ).scalar_one()
        entries = self.db.execute(
            select(SupplierLedgerEntry)
            .where(
                SupplierLedgerEntry.company_id == self.company_id,
                SupplierLedgerEntry.supplier_id == supplier_id,
                SupplierLedgerEntry.document_date >= date_from,
                SupplierLedgerEntry.document_date <= date_to,
            )
            .order_by(SupplierLedgerEntry.document_date)
        ).scalars().all()
        running = money(opening)
        rows = []
        for entry in entries:
            running = money(running + Decimal(entry.credit or 0) - Decimal(entry.debit or 0))
            rows.append(
                {
                    "date": entry.document_date,
                    "document_type": entry.document_type,
                    "document_no": entry.document_no,
                    "due_date": entry.due_date,
                    "debit": str(money(entry.debit)),
                    "credit": str(money(entry.credit)),
                    "balance": str(running),
                }
            )
        return ReportResult(
            code="supplier_statement",
            title=f"Supplier Statement — {supplier.name if supplier else ''}",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "document_type", "title": "Type"},
                {"key": "document_no", "title": "Document"},
                {"key": "due_date", "title": "Due", "type": "date"},
                {"key": "debit", "title": "Debit", "type": "money"},
                {"key": "credit", "title": "Credit", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={"opening_balance": str(money(opening)), "closing_balance": str(running)},
            meta={"supplier_id": supplier_id, "date_from": date_from, "date_to": date_to},
        )

    def receivable_ageing(self, *, as_of: date | None = None) -> ReportResult:
        as_of = as_of or date.today()
        ledger_rows = self.db.execute(
            select(CustomerLedgerEntry, Customer)
            .join(Customer, Customer.id == CustomerLedgerEntry.customer_id)
            .where(
                CustomerLedgerEntry.company_id == self.company_id,
                CustomerLedgerEntry.is_open.is_(True),
            )
        ).all()
        buckets: dict[uuid.UUID, dict[str, Any]] = {}
        for entry, customer in ledger_rows:
            outstanding = money(entry.balance)
            if outstanding <= 0:
                continue
            due = entry.due_date or entry.document_date
            days = (as_of - due).days
            key = "current" if days <= 0 else "0_30" if days <= 30 else "31_60" if days <= 60 else "61_90" if days <= 90 else "over_90"
            record = buckets.setdefault(
                customer.id,
                {
                    "customer_code": customer.code,
                    "customer_name": customer.name,
                    "current": Decimal("0"),
                    "0_30": Decimal("0"),
                    "31_60": Decimal("0"),
                    "61_90": Decimal("0"),
                    "over_90": Decimal("0"),
                    "total": Decimal("0"),
                },
            )
            record[key] += outstanding
            record["total"] += outstanding
        rows = [
            {key: (str(value) if isinstance(value, Decimal) else value) for key, value in record.items()}
            for record in sorted(buckets.values(), key=lambda item: item["total"], reverse=True)
        ]
        totals = {
            bucket: str(money(sum((record[bucket] for record in buckets.values()), Decimal("0"))))
            for bucket in ("current", "0_30", "31_60", "61_90", "over_90", "total")
        }
        return ReportResult(
            code="receivable_ageing",
            title="Accounts Receivable Ageing",
            columns=[
                {"key": "customer_code", "title": "Code"},
                {"key": "customer_name", "title": "Customer"},
                {"key": "current", "title": "Current", "type": "money"},
                {"key": "0_30", "title": "1-30", "type": "money"},
                {"key": "31_60", "title": "31-60", "type": "money"},
                {"key": "61_90", "title": "61-90", "type": "money"},
                {"key": "over_90", "title": "90+", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
            ],
            rows=rows,
            totals=totals,
            meta={"as_of": as_of},
        )

    def payable_ageing(self, *, as_of: date | None = None) -> ReportResult:
        as_of = as_of or date.today()
        ledger_rows = self.db.execute(
            select(SupplierLedgerEntry, Supplier)
            .join(Supplier, Supplier.id == SupplierLedgerEntry.supplier_id)
            .where(SupplierLedgerEntry.company_id == self.company_id, SupplierLedgerEntry.is_open.is_(True))
        ).all()
        buckets: dict[uuid.UUID, dict[str, Any]] = {}
        for entry, supplier in ledger_rows:
            outstanding = money(-entry.balance)
            if outstanding <= 0:
                continue
            due = entry.due_date or entry.document_date
            days = (as_of - due).days
            key = "current" if days <= 0 else "0_30" if days <= 30 else "31_60" if days <= 60 else "61_90" if days <= 90 else "over_90"
            record = buckets.setdefault(
                supplier.id,
                {
                    "supplier_code": supplier.code,
                    "supplier_name": supplier.name,
                    "current": Decimal("0"),
                    "0_30": Decimal("0"),
                    "31_60": Decimal("0"),
                    "61_90": Decimal("0"),
                    "over_90": Decimal("0"),
                    "total": Decimal("0"),
                },
            )
            record[key] += outstanding
            record["total"] += outstanding
        rows = [
            {key: (str(value) if isinstance(value, Decimal) else value) for key, value in record.items()}
            for record in sorted(buckets.values(), key=lambda item: item["total"], reverse=True)
        ]
        totals = {
            bucket: str(money(sum((record[bucket] for record in buckets.values()), Decimal("0"))))
            for bucket in ("current", "0_30", "31_60", "61_90", "over_90", "total")
        }
        return ReportResult(
            code="payable_ageing",
            title="Accounts Payable Ageing",
            columns=[
                {"key": "supplier_code", "title": "Code"},
                {"key": "supplier_name", "title": "Supplier"},
                {"key": "current", "title": "Current", "type": "money"},
                {"key": "0_30", "title": "1-30", "type": "money"},
                {"key": "31_60", "title": "31-60", "type": "money"},
                {"key": "61_90", "title": "61-90", "type": "money"},
                {"key": "over_90", "title": "90+", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
            ],
            rows=rows,
            totals=totals,
            meta={"as_of": as_of},
        )

    # -------------------------------------------------------------- inventory
    def inventory_valuation(
        self, *, warehouse_id: uuid.UUID | None = None, as_of: date | None = None, include_zero: bool = False
    ) -> ReportResult:
        stmt = (
            select(StockBalance, Product, Warehouse)
            .join(Product, Product.id == StockBalance.product_id)
            .join(Warehouse, Warehouse.id == StockBalance.warehouse_id)
            .where(StockBalance.company_id == self.company_id)
            .order_by(Product.sku)
        )
        if warehouse_id:
            stmt = stmt.where(StockBalance.warehouse_id == warehouse_id)
        rows = []
        total_value = Decimal("0")
        for balance, product, warehouse in self.db.execute(stmt).all():
            if not include_zero and (balance.quantity or 0) == 0:
                continue
            total_value += Decimal(balance.total_value or 0)
            rows.append(
                {
                    "sku": product.sku,
                    "product": product.name,
                    "warehouse": warehouse.name,
                    "quantity": str(quantity(balance.quantity)),
                    "average_cost": str(money(balance.average_cost)),
                    "total_value": str(money(balance.total_value)),
                    "last_movement": balance.last_movement_at,
                }
            )
        return ReportResult(
            code="inventory_valuation",
            title="Inventory Valuation",
            columns=[
                {"key": "sku", "title": "SKU"},
                {"key": "product", "title": "Product"},
                {"key": "warehouse", "title": "Warehouse"},
                {"key": "quantity", "title": "Quantity", "type": "number"},
                {"key": "average_cost", "title": "Avg Cost", "type": "money"},
                {"key": "total_value", "title": "Value", "type": "money"},
            ],
            rows=rows,
            totals={"total_value": str(money(total_value)), "lines": len(rows)},
            meta={"as_of": as_of, "warehouse_id": warehouse_id},
        )

    def stock_movement_report(
        self, *, product_id: uuid.UUID | None = None, date_from: date | None = None, date_to: date | None = None, limit: int = 2000
    ) -> ReportResult:
        stmt = (
            select(StockLedgerEntry, Product)
            .join(Product, Product.id == StockLedgerEntry.product_id)
            .where(StockLedgerEntry.company_id == self.company_id)
            .order_by(StockLedgerEntry.entry_date.desc(), StockLedgerEntry.created_at.desc())
            .limit(limit)
        )
        if product_id:
            stmt = stmt.where(StockLedgerEntry.product_id == product_id)
        if date_from:
            stmt = stmt.where(StockLedgerEntry.entry_date >= date_from)
        if date_to:
            stmt = stmt.where(StockLedgerEntry.entry_date <= date_to)
        rows = []
        total_in = total_out = Decimal("0")
        for entry, product in self.db.execute(stmt).all():
            if entry.direction == "in":
                total_in += Decimal(entry.base_quantity or 0)
            else:
                total_out += Decimal(entry.base_quantity or 0)
            rows.append(
                {
                    "date": entry.entry_date,
                    "movement": entry.movement_type,
                    "direction": entry.direction,
                    "sku": product.sku,
                    "product": product.name,
                    "quantity": str(quantity(entry.base_quantity)),
                    "unit_cost": str(money(entry.unit_cost)),
                    "total_cost": str(money(entry.total_cost)),
                    "reference": entry.reference_no,
                    "balance": str(quantity(entry.balance_quantity)),
                }
            )
        return ReportResult(
            code="stock_movement",
            title="Stock Movement",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "movement", "title": "Movement"},
                {"key": "direction", "title": "In/Out"},
                {"key": "sku", "title": "SKU"},
                {"key": "product", "title": "Product"},
                {"key": "quantity", "title": "Qty", "type": "number"},
                {"key": "unit_cost", "title": "Unit Cost", "type": "money"},
                {"key": "total_cost", "title": "Value", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "number"},
                {"key": "reference", "title": "Reference"},
            ],
            rows=rows,
            totals={"total_in": str(quantity(total_in)), "total_out": str(quantity(total_out))},
            meta={"date_from": date_from, "date_to": date_to},
        )

    def reorder_report(self, *, warehouse_id: uuid.UUID | None = None) -> ReportResult:
        from app.services.inventory_service import InventoryService

        items = InventoryService(self.db, self.company_id).low_stock_items(warehouse_id, limit=500)
        rows = [
            {
                "sku": item["sku"],
                "product": item["name"],
                "quantity": str(quantity(item["quantity"])),
                "reorder_level": str(quantity(item["reorder_level"])),
                "shortage": str(quantity(item["shortage"])),
                "average_cost": str(money(item["average_cost"])),
                "value": str(money(item["value"])),
            }
            for item in items
        ]
        return ReportResult(
            code="reorder_report",
            title="Reorder / Low Stock",
            columns=[
                {"key": "sku", "title": "SKU"},
                {"key": "product", "title": "Product"},
                {"key": "quantity", "title": "On Hand", "type": "number"},
                {"key": "reorder_level", "title": "Reorder Level", "type": "number"},
                {"key": "shortage", "title": "Shortage", "type": "number"},
                {"key": "value", "title": "Stock Value", "type": "money"},
            ],
            rows=rows,
            totals={"items_to_reorder": len(rows)},
        )

    # ------------------------------------------------------------------ sales
    def sales_register(
        self,
        *,
        date_from: date,
        date_to: date,
        customer_id: uuid.UUID | None = None,
        salesperson_id: uuid.UUID | None = None,
        group_by: str = "document",
    ) -> ReportResult:
        stmt = (
            select(SalesInvoice, Customer)
            .join(Customer, Customer.id == SalesInvoice.customer_id)
            .where(
                SalesInvoice.company_id == self.company_id,
                SalesInvoice.deleted_at.is_(None),
                SalesInvoice.document_date >= date_from,
                SalesInvoice.document_date <= date_to,
            )
            .order_by(SalesInvoice.document_date)
        )
        if customer_id:
            stmt = stmt.where(SalesInvoice.customer_id == customer_id)
        if salesperson_id:
            stmt = stmt.where(SalesInvoice.salesperson_id == salesperson_id)
        rows = []
        total_net = total_tax = total_total = Decimal("0")
        for invoice, customer in self.db.execute(stmt).all():
            total_net += Decimal(invoice.subtotal or 0)
            total_tax += Decimal(invoice.tax_amount or 0)
            total_total += Decimal(invoice.total_amount or 0)
            rows.append(
                {
                    "date": invoice.document_date,
                    "document_no": invoice.document_no,
                    "customer": customer.name,
                    "status": invoice.status,
                    "subtotal": str(money(invoice.subtotal)),
                    "discount": str(money(invoice.discount_amount)),
                    "tax": str(money(invoice.tax_amount)),
                    "total": str(money(invoice.total_amount)),
                    "paid": str(money(invoice.paid_amount)),
                    "balance": str(money(invoice.balance_amount)),
                    "salesperson_id": invoice.salesperson_id,
                }
            )
        return ReportResult(
            code="sales_register",
            title="Sales Register",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "document_no", "title": "Invoice"},
                {"key": "customer", "title": "Customer"},
                {"key": "status", "title": "Status"},
                {"key": "subtotal", "title": "Net", "type": "money"},
                {"key": "discount", "title": "Discount", "type": "money"},
                {"key": "tax", "title": "Tax", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
                {"key": "paid", "title": "Paid", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={"net": str(money(total_net)), "tax": str(money(total_tax)), "total": str(money(total_total))},
            meta={"date_from": date_from, "date_to": date_to, "group_by": group_by},
        )

    def sales_by_product(self, *, date_from: date, date_to: date, limit: int = 100) -> ReportResult:
        stmt = (
            select(
                Product.sku,
                Product.name,
                func.sum(SalesInvoiceLine.quantity),
                func.sum(SalesInvoiceLine.net_amount),
                func.sum(SalesInvoiceLine.line_total),
                func.count(func.distinct(SalesInvoiceLine.invoice_id)),
            )
            .join(SalesInvoice, SalesInvoice.id == SalesInvoiceLine.invoice_id)
            .join(Product, Product.id == SalesInvoiceLine.product_id)
            .where(
                SalesInvoiceLine.company_id == self.company_id,
                SalesInvoice.status == DocumentStatus.POSTED.value,
                SalesInvoice.document_date >= date_from,
                SalesInvoice.document_date <= date_to,
            )
            .group_by(Product.sku, Product.name)
            .order_by(func.sum(SalesInvoiceLine.line_total).desc())
            .limit(limit)
        )
        rows = []
        total = Decimal("0")
        for sku, name, qty_value, net, gross, invoice_count in self.db.execute(stmt).all():
            total += Decimal(gross or 0)
            rows.append(
                {
                    "sku": sku,
                    "product": name,
                    "quantity": str(quantity(qty_value)),
                    "net": str(money(net)),
                    "total": str(money(gross)),
                    "invoice_count": int(invoice_count or 0),
                }
            )
        return ReportResult(
            code="sales_by_product",
            title="Sales by Product",
            columns=[
                {"key": "sku", "title": "SKU"},
                {"key": "product", "title": "Product"},
                {"key": "quantity", "title": "Quantity", "type": "number"},
                {"key": "net", "title": "Net", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
                {"key": "invoice_count", "title": "Invoices", "type": "number"},
            ],
            rows=rows,
            totals={"total": str(money(total))},
            meta={"date_from": date_from, "date_to": date_to},
        )

    def sales_by_salesperson(self, *, date_from: date, date_to: date) -> ReportResult:
        from app.models.identity import User

        stmt = (
            select(
                User.full_name,
                func.count(func.distinct(SalesInvoice.id)),
                func.sum(SalesInvoice.total_amount),
                func.sum(SalesInvoice.commission_amount),
            )
            .join(SalesInvoice, SalesInvoice.salesperson_id == User.id)
            .where(
                SalesInvoice.company_id == self.company_id,
                SalesInvoice.document_date >= date_from,
                SalesInvoice.document_date <= date_to,
                SalesInvoice.is_return.is_(False),
            )
            .group_by(User.full_name)
            .order_by(func.sum(SalesInvoice.total_amount).desc())
        )
        rows = []
        total = Decimal("0")
        for name, count_value, amount, commission in self.db.execute(stmt).all():
            total += Decimal(amount or 0)
            rows.append(
                {
                    "salesperson": name,
                    "invoices": int(count_value or 0),
                    "sales": str(money(amount)),
                    "commission": str(money(commission)),
                }
            )
        return ReportResult(
            code="sales_by_salesperson",
            title="Sales by Salesperson",
            columns=[
                {"key": "salesperson", "title": "Salesperson"},
                {"key": "invoices", "title": "Invoices", "type": "number"},
                {"key": "sales", "title": "Sales", "type": "money"},
                {"key": "commission", "title": "Commission", "type": "money"},
            ],
            rows=rows,
            totals={"total_sales": str(money(total))},
            meta={"date_from": date_from, "date_to": date_to},
        )

    # -------------------------------------------------------------- purchases
    def purchase_register(
        self, *, date_from: date, date_to: date, supplier_id: uuid.UUID | None = None
    ) -> ReportResult:
        stmt = (
            select(PurchaseInvoice, Supplier)
            .join(Supplier, Supplier.id == PurchaseInvoice.supplier_id)
            .where(
                PurchaseInvoice.company_id == self.company_id,
                PurchaseInvoice.deleted_at.is_(None),
                PurchaseInvoice.document_date >= date_from,
                PurchaseInvoice.document_date <= date_to,
            )
            .order_by(PurchaseInvoice.document_date)
        )
        if supplier_id:
            stmt = stmt.where(PurchaseInvoice.supplier_id == supplier_id)
        rows = []
        total = Decimal("0")
        for invoice, supplier in self.db.execute(stmt).all():
            total += Decimal(invoice.total_amount or 0)
            rows.append(
                {
                    "date": invoice.document_date,
                    "document_no": invoice.document_no,
                    "supplier_invoice_no": invoice.supplier_invoice_no,
                    "supplier": supplier.name,
                    "status": invoice.status,
                    "subtotal": str(money(invoice.subtotal)),
                    "tax": str(money(invoice.tax_amount)),
                    "total": str(money(invoice.total_amount)),
                    "paid": str(money(invoice.paid_amount)),
                    "balance": str(money(invoice.balance_amount)),
                }
            )
        return ReportResult(
            code="purchase_register",
            title="Purchase Register",
            columns=[
                {"key": "date", "title": "Date", "type": "date"},
                {"key": "document_no", "title": "Document"},
                {"key": "supplier_invoice_no", "title": "Supplier Ref"},
                {"key": "supplier", "title": "Supplier"},
                {"key": "status", "title": "Status"},
                {"key": "subtotal", "title": "Net", "type": "money"},
                {"key": "tax", "title": "Tax", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
                {"key": "balance", "title": "Balance", "type": "money"},
            ],
            rows=rows,
            totals={"total": str(money(total))},
            meta={"date_from": date_from, "date_to": date_to},
        )

    def purchase_by_supplier(self, *, date_from: date, date_to: date) -> ReportResult:
        stmt = (
            select(
                Supplier.code,
                Supplier.name,
                func.count(func.distinct(PurchaseInvoice.id)),
                func.sum(PurchaseInvoice.total_amount),
            )
            .join(PurchaseInvoice, PurchaseInvoice.supplier_id == Supplier.id)
            .where(
                PurchaseInvoice.company_id == self.company_id,
                PurchaseInvoice.document_date >= date_from,
                PurchaseInvoice.document_date <= date_to,
            )
            .group_by(Supplier.code, Supplier.name)
            .order_by(func.sum(PurchaseInvoice.total_amount).desc())
        )
        rows = []
        total = Decimal("0")
        for code, name, count_value, amount in self.db.execute(stmt).all():
            total += Decimal(amount or 0)
            rows.append(
                {
                    "supplier_code": code,
                    "supplier": name,
                    "invoices": int(count_value or 0),
                    "purchases": str(money(amount)),
                }
            )
        return ReportResult(
            code="purchase_by_supplier",
            title="Purchases by Supplier",
            columns=[
                {"key": "supplier_code", "title": "Code"},
                {"key": "supplier", "title": "Supplier"},
                {"key": "invoices", "title": "Invoices", "type": "number"},
                {"key": "purchases", "title": "Purchases", "type": "money"},
            ],
            rows=rows,
            totals={"total": str(money(total))},
            meta={"date_from": date_from, "date_to": date_to},
        )

    def supplier_price_comparison(self, *, product_id: uuid.UUID, limit: int = 20) -> ReportResult:
        stmt = (
            select(
                Supplier.code,
                Supplier.name,
                func.min(text("unit_price")),
                func.max(text("unit_price")),
                func.avg(text("unit_price")),
                func.count(),
                func.max(text("price_date")),
            )
            .select_from(Supplier)
            .join(text("supplier_price_history"), text("supplier_price_history.supplier_id = suppliers.id"))
            .where(
                text("supplier_price_history.company_id = :company_id"),
                text("supplier_price_history.product_id = :product_id"),
            )
            .params(company_id=str(self.company_id), product_id=str(product_id))
            .group_by(Supplier.code, Supplier.name)
            .order_by(func.avg(text("unit_price")).asc())
            .limit(limit)
        )
        rows = [
            {
                "supplier_code": code,
                "supplier": name,
                "min_price": str(money(minimum)),
                "max_price": str(money(maximum)),
                "avg_price": str(money(average)),
                "quotes": int(count_value or 0),
                "last_price_date": last_date,
            }
            for code, name, minimum, maximum, average, count_value, last_date in self.db.execute(stmt).all()
        ]
        return ReportResult(
            code="supplier_price_comparison",
            title="Supplier Price Comparison",
            columns=[
                {"key": "supplier_code", "title": "Code"},
                {"key": "supplier", "title": "Supplier"},
                {"key": "min_price", "title": "Min", "type": "money"},
                {"key": "max_price", "title": "Max", "type": "money"},
                {"key": "avg_price", "title": "Average", "type": "money"},
                {"key": "quotes", "title": "Quotes", "type": "number"},
                {"key": "last_price_date", "title": "Last Price", "type": "date"},
            ],
            rows=rows,
            totals={"suppliers": len(rows)},
            meta={"product_id": product_id},
        )

    # -------------------------------------------------------------------- tax
    def tax_report(self, *, date_from: date, date_to: date) -> ReportResult:
        sales_rows = self.db.execute(
            select(
                func.coalesce(func.sum(SalesInvoice.subtotal), 0),
                func.coalesce(func.sum(SalesInvoice.tax_amount), 0),
                func.coalesce(func.sum(SalesInvoice.total_amount), 0),
            ).where(
                SalesInvoice.company_id == self.company_id,
                SalesInvoice.status == DocumentStatus.POSTED.value,
                SalesInvoice.document_date >= date_from,
                SalesInvoice.document_date <= date_to,
            )
        ).one()
        purchase_rows = self.db.execute(
            select(
                func.coalesce(func.sum(PurchaseInvoice.subtotal), 0),
                func.coalesce(func.sum(PurchaseInvoice.tax_amount), 0),
                func.coalesce(func.sum(PurchaseInvoice.total_amount), 0),
            ).where(
                PurchaseInvoice.company_id == self.company_id,
                PurchaseInvoice.status == DocumentStatus.POSTED.value,
                PurchaseInvoice.document_date >= date_from,
                PurchaseInvoice.document_date <= date_to,
            )
        ).one()
        output_tax = money(sales_rows[1])
        input_tax = money(purchase_rows[1])
        rows = [
            {"section": "Sales", "base": str(money(sales_rows[0])), "tax": str(output_tax), "total": str(money(sales_rows[2]))},
            {
                "section": "Purchases",
                "base": str(money(purchase_rows[0])),
                "tax": str(input_tax),
                "total": str(money(purchase_rows[2])),
            },
        ]
        return ReportResult(
            code="tax_report",
            title="Tax Report",
            columns=[
                {"key": "section", "title": "Section"},
                {"key": "base", "title": "Taxable Base", "type": "money"},
                {"key": "tax", "title": "Tax", "type": "money"},
                {"key": "total", "title": "Total", "type": "money"},
            ],
            rows=rows,
            totals={
                "output_tax": str(output_tax),
                "input_tax": str(input_tax),
                "net_tax_payable": str(money(output_tax - input_tax)),
            },
            meta={"date_from": date_from, "date_to": date_to},
        )

    # --------------------------------------------------------------------- HR
    def payroll_summary(self, *, payroll_run_id: uuid.UUID) -> ReportResult:
        run = self.db.get(PayrollRun, payroll_run_id)
        if run is None:
            raise BusinessRuleError("Payroll run not found")
        payslips = self.db.execute(
            select(Payslip, Employee)
            .join(Employee, Employee.id == Payslip.employee_id)
            .where(Payslip.payroll_run_id == payroll_run_id)
        ).all()
        rows = [
            {
                "employee_no": employee.employee_no,
                "employee": employee.full_name,
                "basic": str(money(payslip.basic_salary)),
                "allowances": str(
                    money(
                        Decimal(payslip.housing_allowance or 0)
                        + Decimal(payslip.transport_allowance or 0)
                        + Decimal(payslip.other_allowances or 0)
                    )
                ),
                "overtime": str(money(payslip.overtime_amount)),
                "gross": str(money(payslip.gross_pay)),
                "deductions": str(money(payslip.total_deductions)),
                "net": str(money(payslip.net_pay)),
            }
            for payslip, employee in payslips
        ]
        return ReportResult(
            code="payroll_summary",
            title=f"Payroll Summary — {run.run_no}",
            columns=[
                {"key": "employee_no", "title": "Employee No"},
                {"key": "employee", "title": "Employee"},
                {"key": "basic", "title": "Basic", "type": "money"},
                {"key": "allowances", "title": "Allowances", "type": "money"},
                {"key": "overtime", "title": "Overtime", "type": "money"},
                {"key": "gross", "title": "Gross", "type": "money"},
                {"key": "deductions", "title": "Deductions", "type": "money"},
                {"key": "net", "title": "Net", "type": "money"},
            ],
            rows=rows,
            totals={
                "employees": len(rows),
                "gross": str(money(run.total_gross)),
                "deductions": str(money(run.total_deductions)),
                "net": str(money(run.total_net)),
            },
            meta={"payroll_run_id": payroll_run_id},
        )

    def attendance_summary(self, *, date_from: date, date_to: date, employee_id: uuid.UUID | None = None) -> ReportResult:
        stmt = (
            select(
                Employee.employee_no,
                Employee.first_name,
                Employee.last_name,
                func.count(AttendanceRecord.id),
                func.coalesce(func.sum(AttendanceRecord.overtime_minutes), 0),
                func.coalesce(func.sum(AttendanceRecord.late_minutes), 0),
            )
            .join(AttendanceRecord, AttendanceRecord.employee_id == Employee.id)
            .where(
                AttendanceRecord.company_id == self.company_id,
                AttendanceRecord.attendance_date >= date_from,
                AttendanceRecord.attendance_date <= date_to,
            )
            .group_by(Employee.employee_no, Employee.first_name, Employee.last_name)
        )
        if employee_id:
            stmt = stmt.where(Employee.id == employee_id)
        rows = [
            {
                "employee_no": number,
                "employee": f"{first} {last}",
                "days": int(days or 0),
                "overtime_hours": str(quantity(Decimal(overtime or 0) / Decimal("60"))),
                "late_hours": str(quantity(Decimal(late or 0) / Decimal("60"))),
            }
            for number, first, last, days, overtime, late in self.db.execute(stmt).all()
        ]
        return ReportResult(
            code="attendance_summary",
            title="Attendance Summary",
            columns=[
                {"key": "employee_no", "title": "Employee No"},
                {"key": "employee", "title": "Employee"},
                {"key": "days", "title": "Present Days", "type": "number"},
                {"key": "overtime_hours", "title": "Overtime (h)", "type": "number"},
                {"key": "late_hours", "title": "Late (h)", "type": "number"},
            ],
            rows=rows,
            totals={"employees": len(rows)},
            meta={"date_from": date_from, "date_to": date_to},
        )

    def leave_summary(self, *, year: int) -> ReportResult:
        stmt = (
            select(
                Employee.employee_no,
                Employee.first_name,
                Employee.last_name,
                func.count(LeaveRequest.id),
                func.coalesce(func.sum(LeaveRequest.total_days), 0),
            )
            .join(LeaveRequest, LeaveRequest.employee_id == Employee.id)
            .where(
                LeaveRequest.company_id == self.company_id,
                LeaveRequest.start_date >= date(year, 1, 1),
                LeaveRequest.start_date <= date(year, 12, 31),
            )
            .group_by(Employee.employee_no, Employee.first_name, Employee.last_name)
        )
        rows = [
            {
                "employee_no": number,
                "employee": f"{first} {last}",
                "requests": int(count_value or 0),
                "days": str(quantity(days)),
            }
            for number, first, last, count_value, days in self.db.execute(stmt).all()
        ]
        return ReportResult(
            code="leave_summary",
            title="Leave Summary",
            columns=[
                {"key": "employee_no", "title": "Employee No"},
                {"key": "employee", "title": "Employee"},
                {"key": "requests", "title": "Requests", "type": "number"},
                {"key": "days", "title": "Days", "type": "number"},
            ],
            rows=rows,
            totals={"employees": len(rows)},
            meta={"year": year},
        )

    # ---------------------------------------------------------------- projects
    def project_profitability(self, *, project_id: uuid.UUID | None = None) -> ReportResult:
        stmt = select(Project).where(Project.company_id == self.company_id, Project.deleted_at.is_(None))
        if project_id:
            stmt = stmt.where(Project.id == project_id)
        rows = []
        total_contract = total_cost = total_profit = Decimal("0")
        for project in self.db.execute(stmt.order_by(Project.project_no)).scalars().all():
            cost = Decimal(project.total_actual_cost)
            revenue = Decimal(project.invoiced_amount or 0)
            profit = money(revenue - cost)
            margin = (profit / revenue * Decimal("100")).quantize(Decimal("0.01")) if revenue else Decimal("0")
            total_contract += Decimal(project.contract_value or 0)
            total_cost += cost
            total_profit += profit
            rows.append(
                {
                    "project_no": project.project_no,
                    "name": project.name,
                    "status": project.status,
                    "contract_value": str(money(project.contract_value)),
                    "budget": str(money(project.budget_amount)),
                    "actual_cost": str(money(cost)),
                    "invoiced": str(money(revenue)),
                    "profit": str(profit),
                    "margin_percent": str(margin),
                    "progress": str(project.progress_percent or 0),
                }
            )
        return ReportResult(
            code="project_profitability",
            title="Project Profitability",
            columns=[
                {"key": "project_no", "title": "Project"},
                {"key": "name", "title": "Name"},
                {"key": "status", "title": "Status"},
                {"key": "contract_value", "title": "Contract", "type": "money"},
                {"key": "budget", "title": "Budget", "type": "money"},
                {"key": "actual_cost", "title": "Actual Cost", "type": "money"},
                {"key": "invoiced", "title": "Invoiced", "type": "money"},
                {"key": "profit", "title": "Profit", "type": "money"},
                {"key": "margin_percent", "title": "Margin %", "type": "number"},
            ],
            rows=rows,
            totals={"contract_value": str(money(total_contract)), "cost": str(money(total_cost)), "profit": str(money(total_profit))},
            meta={"project_id": project_id},
        )

    def timesheet_report(self, *, date_from: date, date_to: date, project_id: uuid.UUID | None = None) -> ReportResult:
        """Aggregated timesheet hours / billable value per employee and project."""
        detail = self.db.execute(
            select(TimesheetLine, text("employees.employee_no"), text("employees.first_name"), text("employees.last_name"))
            .select_from(TimesheetLine)
            .join(text("timesheets"), text("timesheets.id = timesheet_lines.timesheet_id"))
            .join(text("employees"), text("employees.id = timesheets.employee_id"))
            .where(
                TimesheetLine.company_id == self.company_id,
                TimesheetLine.work_date >= date_from,
                TimesheetLine.work_date <= date_to,
            )
        ).all()
        aggregated: dict[tuple, dict[str, Any]] = {}
        for line, employee_no, first_name, last_name in detail:
            key = (employee_no, line.project_id)
            record = aggregated.setdefault(
                key,
                {
                    "employee_no": employee_no,
                    "employee": f"{first_name} {last_name}",
                    "project_id": line.project_id,
                    "hours": Decimal("0"),
                    "billable": Decimal("0"),
                    "cost": Decimal("0"),
                },
            )
            record["hours"] += Decimal(line.hours or 0)
            record["billable"] += Decimal(line.billable_amount or 0)
            record["cost"] += Decimal(line.cost_amount or 0)
        rows = []
        for record in aggregated.values():
            project = self.db.get(Project, record["project_id"]) if record["project_id"] else None
            rows.append(
                {
                    "employee_no": record["employee_no"],
                    "employee": record["employee"],
                    "project": project.project_no if project else "—",
                    "hours": str(quantity(record["hours"])),
                    "billable": str(money(record["billable"])),
                    "cost": str(money(record["cost"])),
                }
            )
        return ReportResult(
            code="timesheet_report",
            title="Timesheet Report",
            columns=[
                {"key": "employee_no", "title": "Employee No"},
                {"key": "employee", "title": "Employee"},
                {"key": "project", "title": "Project"},
                {"key": "hours", "title": "Hours", "type": "number"},
                {"key": "billable", "title": "Billable", "type": "money"},
                {"key": "cost", "title": "Cost", "type": "money"},
            ],
            rows=rows,
            totals={"entries": len(rows)},
            meta={"date_from": date_from, "date_to": date_to},
        )

    # -------------------------------------------------------------- registers
    def available_reports(self) -> list[dict[str, Any]]:
        return [
            {"code": "trial_balance", "title": "Trial Balance", "module": "accounting"},
            {"code": "general_ledger", "title": "General Ledger", "module": "accounting"},
            {"code": "income_statement", "title": "Profit & Loss", "module": "accounting"},
            {"code": "balance_sheet", "title": "Balance Sheet", "module": "accounting"},
            {"code": "cash_flow", "title": "Cash Flow", "module": "accounting"},
            {"code": "receivable_ageing", "title": "AR Ageing", "module": "accounting"},
            {"code": "payable_ageing", "title": "AP Ageing", "module": "accounting"},
            {"code": "customer_statement", "title": "Customer Statement", "module": "accounting"},
            {"code": "supplier_statement", "title": "Supplier Statement", "module": "accounting"},
            {"code": "tax_report", "title": "Tax Report", "module": "tax"},
            {"code": "inventory_valuation", "title": "Inventory Valuation", "module": "inventory"},
            {"code": "stock_movement", "title": "Stock Movement", "module": "inventory"},
            {"code": "reorder_report", "title": "Reorder / Low Stock", "module": "inventory"},
            {"code": "sales_register", "title": "Sales Register", "module": "sales"},
            {"code": "sales_by_product", "title": "Sales by Product", "module": "sales"},
            {"code": "sales_by_salesperson", "title": "Sales by Salesperson", "module": "sales"},
            {"code": "purchase_register", "title": "Purchase Register", "module": "purchasing"},
            {"code": "purchase_by_supplier", "title": "Purchases by Supplier", "module": "purchasing"},
            {"code": "supplier_price_comparison", "title": "Supplier Price Comparison", "module": "purchasing"},
            {"code": "payroll_summary", "title": "Payroll Summary", "module": "hr"},
            {"code": "attendance_summary", "title": "Attendance Summary", "module": "hr"},
            {"code": "leave_summary", "title": "Leave Summary", "module": "hr"},
            {"code": "project_profitability", "title": "Project Profitability", "module": "projects"},
            {"code": "timesheet_report", "title": "Timesheet Report", "module": "projects"},
        ]

    def run(self, code: str, params: dict[str, Any]) -> ReportResult:
        """Dispatch a report by code (used by the API and the scheduler)."""
        handlers: dict[str, Callable[..., ReportResult]] = {
            "trial_balance": lambda: self.trial_balance(
                date_from=params.get("date_from"),
                date_to=params.get("date_to"),
                include_zero=bool(params.get("include_zero", False)),
                branch_id=params.get("branch_id"),
            ),
            "general_ledger": lambda: self.general_ledger(
                account_id=params.get("account_id"),
                date_from=params.get("date_from"),
                date_to=params.get("date_to"),
                party_id=params.get("party_id"),
            ),
            "income_statement": lambda: self.income_statement(
                date_from=params["date_from"],
                date_to=params["date_to"],
                branch_id=params.get("branch_id"),
                compare=bool(params.get("compare", False)),
            ),
            "balance_sheet": lambda: self.balance_sheet(
                as_of=params.get("as_of") or params.get("date_to") or date.today(), branch_id=params.get("branch_id")
            ),
            "cash_flow": lambda: self.cash_flow(date_from=params["date_from"], date_to=params["date_to"]),
            "receivable_ageing": lambda: self.receivable_ageing(as_of=params.get("as_of")),
            "payable_ageing": lambda: self.payable_ageing(as_of=params.get("as_of")),
            "customer_statement": lambda: self.customer_statement(
                customer_id=params["customer_id"],
                date_from=params.get("date_from") or date(date.today().year, 1, 1),
                date_to=params.get("date_to") or date.today(),
            ),
            "supplier_statement": lambda: self.supplier_statement(
                supplier_id=params["supplier_id"],
                date_from=params.get("date_from") or date(date.today().year, 1, 1),
                date_to=params.get("date_to") or date.today(),
            ),
            "tax_report": lambda: self.tax_report(date_from=params["date_from"], date_to=params["date_to"]),
            "inventory_valuation": lambda: self.inventory_valuation(
                warehouse_id=params.get("warehouse_id"),
                as_of=params.get("as_of"),
                include_zero=bool(params.get("include_zero", False)),
            ),
            "stock_movement": lambda: self.stock_movement_report(
                product_id=params.get("product_id"),
                date_from=params.get("date_from"),
                date_to=params.get("date_to"),
            ),
            "reorder_report": lambda: self.reorder_report(warehouse_id=params.get("warehouse_id")),
            "sales_register": lambda: self.sales_register(
                date_from=params["date_from"],
                date_to=params["date_to"],
                customer_id=params.get("customer_id"),
                salesperson_id=params.get("salesperson_id"),
                group_by=params.get("group_by", "document"),
            ),
            "sales_by_product": lambda: self.sales_by_product(
                date_from=params["date_from"], date_to=params["date_to"]
            ),
            "sales_by_salesperson": lambda: self.sales_by_salesperson(
                date_from=params["date_from"], date_to=params["date_to"]
            ),
            "purchase_register": lambda: self.purchase_register(
                date_from=params["date_from"], date_to=params["date_to"], supplier_id=params.get("supplier_id")
            ),
            "purchase_by_supplier": lambda: self.purchase_by_supplier(
                date_from=params["date_from"], date_to=params["date_to"]
            ),
            "supplier_price_comparison": lambda: self.supplier_price_comparison(product_id=params["product_id"]),
            "payroll_summary": lambda: self.payroll_summary(payroll_run_id=params["payroll_run_id"]),
            "attendance_summary": lambda: self.attendance_summary(
                date_from=params["date_from"], date_to=params["date_to"], employee_id=params.get("employee_id")
            ),
            "leave_summary": lambda: self.leave_summary(year=int(params.get("year") or date.today().year)),
            "project_profitability": lambda: self.project_profitability(project_id=params.get("project_id")),
            "timesheet_report": lambda: self.timesheet_report(
                date_from=params["date_from"], date_to=params["date_to"], project_id=params.get("project_id")
            ),
        }
        if code not in handlers:
            raise ValidationFailure(f"Unknown report '{code}'", available=sorted(handlers))
        try:
            return handlers[code]()
        except KeyError as exc:
            raise ValidationFailure(f"Missing required report parameter: {exc.args[0]}") from exc
